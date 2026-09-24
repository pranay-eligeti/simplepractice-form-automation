"""Excel input adapter using openpyxl."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def read_records(path: str | Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
        return [dict(zip(headers, row)) for row in rows if any(value is not None for value in row)]
    finally:
        workbook.close()
